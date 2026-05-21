import openpyxl
from openpyxl.utils import get_column_letter

file_path = r'c:\Users\emedrano\Downloads\Lab\Medrano Zamora.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)
ws = wb['Fundamento del Pronóstico']

print("=== VERIFYING CELL REFERENCES ===\n")

# Check CUADRO 7 structure (rows 158-200)
print("CUADRO 7: COSTOS DE OPERACION (Rows 158-200)")
print("Checking rows 165-190 for reference values:\n")

for row in range(165, 191):
    row_data = []
    for col in range(1, 8):
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v is not None and v != '':
            if isinstance(v, str) and v.startswith('='):
                row_data.append(f"{get_column_letter(col)}{row}={v[:30]}")
            else:
                row_data.append(f"{get_column_letter(col)}{row}:{v}")
    if row_data:
        print(f"Row {row}: {' | '.join(row_data[:4])}")

print("\n\n=== CUADRO 3: PROGRAMA DE PRODUCCION ===")
for row in range(83, 90):
    row_data = []
    for col in range(1, 8):
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v is not None and v != '':
            if isinstance(v, str) and v.startswith('='):
                row_data.append(f"{get_column_letter(col)}{row}={v[:30]}")
            else:
                row_data.append(f"{get_column_letter(col)}{row}:{v}")
    if row_data:
        print(f"Row {row}: {' | '.join(row_data)}")

wb.close()
