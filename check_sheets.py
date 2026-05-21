import openpyxl

file_path = r'c:\Users\emedrano\Downloads\Lab\Medrano Zamora.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)

print("=== SHEETS IN WORKBOOK ===")
for name in wb.sheetnames:
    print(f"  - {name}")

# Check each sheet for CUADRO
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== Sheet: {sheet_name} ===")
    
    found = []
    for row in range(1, 300):
        for col in range(1, 20):
            cell = ws.cell(row=row, column=col)
            val = cell.value
            if val and isinstance(val, str) and 'CUADRO' in val.upper():
                found.append((cell.coordinate, row, val))
                
    if found:
        print(f"Found {len(found)} CUADRO references:")
        for coord, row, val in found[:20]:  # Limit output
            print(f"  {coord} (Row {row}): {val}")
    else:
        print("  No CUADRO references found")

wb.close()
