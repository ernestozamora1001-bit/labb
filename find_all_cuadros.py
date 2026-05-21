import openpyxl

file_path = r'c:\Users\emedrano\Downloads\Lab\Medrano Zamora.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)
ws = wb.active

print("=== FINDING ALL CUADROS ===\n")

cuadros = []
for row in range(1, 250):
    for col in range(1, 15):
        cell = ws.cell(row=row, column=col)
        val = cell.value
        if val and isinstance(val, str) and 'CUADRO' in val.upper() and len(val) < 50:
            cuadros.append((row, col, cell.coordinate, val))
            break

print("CUADROS found:")
for row, col, coord, val in cuadros:
    print(f"  Row {row} (Col {col}, {coord}): {val}")

# Show context around each CUADRO
print("\n\n=== DETAILED VIEW OF EACH CUADRO ===")
for i, (row, col, coord, val) in enumerate(cuadros):
    end_row = cuadros[i+1][0] if i+1 < len(cuadros) else min(row + 30, 250)
    print(f"\n{val} (Rows {row}-{end_row-1}):")
    for r in range(row, min(row+15, end_row)):
        row_data = []
        for c in range(1, 10):
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if v is not None and v != '':
                row_data.append(f"[{c}]{v}")
        if row_data:
            print(f"  Row {r}: {' | '.join(row_data[:6])}")  # Limit output

wb.close()
