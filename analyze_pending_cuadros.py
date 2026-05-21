import openpyxl

file_path = r'c:\Users\emedrano\Downloads\Lab\Medrano Zamora.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)
ws = wb['Fundamento del Pronóstico']

print("=== DETAILED ANALYSIS OF PENDING CUADROS ===\n")

# CUADRO 4: Rows 91-106
print("=" * 60)
print("CUADRO 4: REQUERIMIENTO DE MATERIALES Y MANO DE OBRA (Rows 91-106)")
print("=" * 60)
for row in range(91, 107):
    print(f"\nRow {row}:")
    for col in range(1, 12):
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v is not None and v != '':
            if isinstance(v, str) and v.startswith('='):
                print(f"  {cell.coordinate}: FORMULA: {v[:50]}")
            else:
                print(f"  {cell.coordinate}: {v}")

print("\n\n" + "=" * 60)
print("CUADRO 5: REQUERIMIENTOS DE CAPITAL DE TRABAJO (Rows 107-135)")
print("=" * 60)
for row in range(107, 136):
    print(f"\nRow {row}:")
    for col in range(1, 12):
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v is not None and v != '':
            if isinstance(v, str) and v.startswith('='):
                print(f"  {cell.coordinate}: FORMULA: {v[:50]}")
            else:
                print(f"  {cell.coordinate}: {v}")

print("\n\n" + "=" * 60)
print("CUADRO 6: COSTOS FIJOS Y VARIABLES (Rows 136-157)")
print("=" * 60)
for row in range(136, 158):
    print(f"\nRow {row}:")
    for col in range(1, 8):
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v is not None and v != '':
            if isinstance(v, str) and v.startswith('='):
                print(f"  {cell.coordinate}: FORMULA: {v[:50]}")
            else:
                print(f"  {cell.coordinate}: {v}")

wb.close()
