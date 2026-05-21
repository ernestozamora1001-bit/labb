import pandas as pd
import openpyxl

file_path = r'c:\Users\emedrano\Downloads\Lab\Medrano Zamora.xlsx'

# Read with pandas
xl = pd.ExcelFile(file_path)
df = xl.parse(xl.sheet_names[0], header=None)

print("=== EXCEL ANALYSIS ===")
print(f"Total rows: {len(df)}")
print(f"Total columns: {len(df.columns)}")

# Find CUADRO sections
cuadro_rows = []
for idx, row in df.iterrows():
    for col in row.index:
        val = str(row[col]) if pd.notna(row[col]) else ""
        if 'CUADRO' in val and len(val) < 40:
            cuadro_rows.append((idx, col, val))
            break

print("\n=== CUADROS FOUND ===")
for row_idx, col_idx, val in cuadro_rows:
    print(f"Row {row_idx}: {val}")

# Read with openpyxl to check formulas
wb = openpyxl.load_workbook(file_path, data_only=False)
ws = wb.active

print("\n=== CHECKING FOR FORMULAS IN CUADRO 1 (Rows 46-66) ===")
for row in range(46, 67):
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            print(f"Cell {cell.coordinate}: {cell.value}")

print("\n=== CHECKING FOR FORMULAS IN CUADRO 2 (Rows 67-81) ===")
for row in range(67, 82):
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            print(f"Cell {cell.coordinate}: {cell.value}")

print("\n=== CUADRO 2 DATA (Rows 67-81, Cols 1-8) ===")
for row in range(67, 82):
    row_data = []
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        val = cell.value if cell.value is not None else ""
        if val:
            row_data.append(f"[{col}]{val}")
    if row_data:
        print(f"Row {row}: {' | '.join(row_data)}")

wb.close()
