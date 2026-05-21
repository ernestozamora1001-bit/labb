import openpyxl

file_path = r'c:\Users\emedrano\Downloads\Lab\Medrano Zamora.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)

# Check TAB1-3 sheet
print("=== SHEET: TAB1-3 ===")
ws = wb['TAB1-3']
for row in range(1, 50):
    row_data = []
    for col in range(1, 10):
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v is not None and v != '':
            if isinstance(v, str) and v.startswith('='):
                row_data.append(f"[{col}]={v[:25]}")
            else:
                row_data.append(f"[{col}]{v}")
    if row_data:
        print(f"Row {row}: {' | '.join(row_data[:6])}")

print("\n\n=== SHEET: TAB4-11 (first 30 rows) ===")
ws = wb['TAB4-11']
for row in range(1, 31):
    row_data = []
    for col in range(1, 10):
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v is not None and v != '':
            if isinstance(v, str) and v.startswith('='):
                row_data.append(f"[{col}]={v[:25]}")
            else:
                row_data.append(f"[{col}]{v}")
    if row_data:
        print(f"Row {row}: {' | '.join(row_data[:6])}")

wb.close()
