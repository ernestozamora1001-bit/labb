import openpyxl
from openpyxl.utils import get_column_letter

file_path = r'c:\Users\emedrano\Downloads\Lab\Medrano Zamora.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)
ws = wb['Fundamento del Pronóstico']

print("=== VERIFYING ADDED FORMULAS ===\n")

# Verify CUADRO 4 formulas
print("CUADRO 4: Rows 100-105 formulas:")
for row in range(100, 106):
    row_data = []
    for col in range(3, 6):  # C, D, E
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v and isinstance(v, str) and v.startswith('='):
            row_data.append(f"{get_column_letter(col)}{row}:{v}")
        elif v:
            row_data.append(f"{get_column_letter(col)}{row}:{v}")
    if row_data:
        print(f"  Row {row}: {' | '.join(row_data)}")

# Verify CUADRO 5 formulas
print("\nCUADRO 5: Working capital formulas (rows 111-134):")
for row in [111, 114, 115, 116, 117, 118, 120, 121, 126, 127, 129, 131, 132, 134]:
    row_data = []
    for col in [2, 3, 4]:  # B, C, D
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v and isinstance(v, str) and v.startswith('='):
            row_data.append(f"{get_column_letter(col)}{row}={v[:40]}")
        elif v:
            row_data.append(f"{get_column_letter(col)}{row}:{v}")
    if row_data:
        print(f"  Row {row}: {' | '.join(row_data)}")

# Verify CUADRO 6 formulas
print("\nCUADRO 6: Totals formulas:")
for row in [153, 154]:
    row_data = []
    for col in [2, 3, 4]:  # B, C, D
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if v and isinstance(v, str) and v.startswith('='):
            row_data.append(f"{get_column_letter(col)}{row}={v[:40]}")
        elif v:
            row_data.append(f"{get_column_letter(col)}{row}:{v}")
    if row_data:
        print(f"  Row {row}: {' | '.join(row_data)}")

# Also check CUADRO 7 row 190
print("\nCUADRO 7: Row 190 (Gastos distribucion) - filled values:")
for col in [4, 5, 6]:  # D, E, F
    cell = ws.cell(row=190, column=col)
    v = cell.value
    if v:
        print(f"  {get_column_letter(col)}190: {v}")

wb.close()

print("\n=== VERIFICATION COMPLETE ===")
